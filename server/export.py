"""Export helpers: turn a re-teach plan (Markdown) into a small .pptx the lecturer can drop into their deck."""
from __future__ import annotations

import io
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


def _sections(md: str) -> list[tuple[str, list[str]]]:
    out: list[tuple[str, list[str]]] = []
    for line in md.splitlines():
        line = line.rstrip()
        if not line.strip():
            continue
        m = re.match(r"^\*\*(.+?)\*\*\s*$", line.strip())
        if m:
            out.append((m.group(1).strip(), []))
            continue
        if not out:
            out.append(("Re-teach", []))
        text = re.sub(r"^\s*[-*]\s+", "", line)
        text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
        out[-1][1].append(text.strip())
    return out


def reteach_pptx(title: str, md: str) -> io.BytesIO:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    ink, teal, slate = RGBColor(0x16, 0x20, 0x2A), RGBColor(0x0E, 0x7C, 0x66), RGBColor(0x6B, 0x77, 0x85)

    def add_slide(head: str, lines: list[str], kicker: str):
        sl = prs.slides.add_slide(blank)
        tb = sl.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.5)).text_frame
        tb.text = kicker
        tb.paragraphs[0].runs[0].font.size, tb.paragraphs[0].runs[0].font.color.rgb = Pt(16), teal
        tb.paragraphs[0].runs[0].font.bold = True
        hb = sl.shapes.add_textbox(Inches(0.8), Inches(1.0), Inches(11.7), Inches(1.2)).text_frame
        hb.word_wrap = True
        hb.text = head
        hb.paragraphs[0].runs[0].font.size, hb.paragraphs[0].runs[0].font.color.rgb = Pt(36), ink
        hb.paragraphs[0].runs[0].font.bold = True
        body = sl.shapes.add_textbox(Inches(0.8), Inches(2.4), Inches(11.7), Inches(4.5)).text_frame
        body.word_wrap = True
        for i, ln in enumerate(lines):
            p = body.paragraphs[0] if i == 0 else body.add_paragraph()
            p.text = ("• " if len(lines) > 1 else "") + ln
            p.runs[0].font.size, p.runs[0].font.color.rgb = Pt(22), ink
            p.space_after = Pt(10)
        foot = sl.shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(11.7), Inches(0.4)).text_frame
        foot.text = "Drafted from students' own explanations · Utschytl"
        foot.paragraphs[0].runs[0].font.size, foot.paragraphs[0].runs[0].font.color.rgb = Pt(11), slate

    for head, lines in _sections(md) or [("Re-teach", [md])]:
        add_slide(head, lines or ["(no content)"], title)
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf
